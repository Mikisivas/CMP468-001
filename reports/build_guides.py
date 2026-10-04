"""Builds the beginner run guide and the defence Q&A document."""
import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from docx_helpers import Report, _shade

HERE = os.path.dirname(os.path.abspath(__file__))


def steps(r, items):
    for i, it in enumerate(items, 1):
        p = r.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.add_run(f"{i}.  ").bold = True
        parts = it.split("**")
        for j, part in enumerate(parts):
            p.add_run(part).bold = j % 2 == 1


def box(r, title, lines, fill="E2F0D9"):
    t = r.doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    _shade(c, fill)
    c.text = ""
    p = c.paragraphs[0]
    p.add_run(title).bold = True
    for ln in lines:
        q = c.add_paragraph()
        q.paragraph_format.line_spacing = 1.15
        parts = ln.split("**")
        for j, part in enumerate(parts):
            run = q.add_run(part)
            run.bold = j % 2 == 1
            run.font.size = Pt(11)
    r.doc.add_paragraph()


def typed(r, text):
    """Something the reader must type exactly."""
    t = r.doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    _shade(c, "F2F2F2")
    c.text = ""
    run = c.paragraphs[0].add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(11)
    run.bold = True
    r.doc.add_paragraph()


def cover(r, title, sub):
    for _ in range(5):
        r.doc.add_paragraph()
    r.center("CMP 468: COMPUTER SECURITY", 13, True)
    r.center(title, 20, True, 14)
    r.center(sub, 12, False, 30)
    r.center("Project 1: VarsityShield (Monitoring, Backup and Recovery)", 12)
    r.center("Project 2: Zaman Lafiya (GIS Early Warning for Farmer-Herder Conflict)", 12, False, 40)
    r.center("October 2026", 11)
    r.page_break()


def qa(r, items):
    for n, (q, a) in enumerate(items, 1):
        p = r.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"Q{n}. {q}")
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
        p2 = r.doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.6)
        p2.paragraph_format.space_after = Pt(10)
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parts = a.split("**")
        for j, part in enumerate(parts):
            p2.add_run(part).bold = j % 2 == 1


# =================================================================== RUN GUIDE
r = Report()
cover(r, "HOW TO RUN BOTH PROJECTS", "A step-by-step guide for beginners (Windows 10 or 11)")

r.h1("BEFORE YOU START")
r.p("Read this page once. It saves you an hour of confusion.")
r.table("What you need", ["Item", "Details"],
        [["A laptop", "Windows 10 or Windows 11, at least 4 GB RAM, about 1.5 GB free space"],
         ["Internet", "Needed ONLY for Part 1 and the two setup steps (about 300 MB download). "
                      "After setup, the systems run without internet. Only the street map background "
                      "in Project 2 needs internet."],
         ["Time", "About 30 minutes the first time. After that, starting a demo takes 1 minute."],
         ["The zip file", "CMP468_Projects.zip"]],
        widths=[3.5, 12.5])
box(r, "How to read this guide", [
    "Text in a **grey box** is something you type exactly as shown, then press **Enter**.",
    "Text in a **green box** tells you what you should see if everything is working.",
    "Text in a **yellow box** is a warning. Read it carefully.",
    "Do the parts in order. Do not skip Part 1.",
])
r.h2("What is inside the zip")
r.table("Folders and files", ["Folder or file", "What it is"],
        [["project1-varsityshield", "Project 1 program. You will double-click files whose names start with 1_, 2_, 3_, 4_"],
         ["project2-zamanlafiya", "Project 2 program. Same idea: files starting with 1_, 2_, 3_, 4_"],
         ["reports", "The two project reports (Word documents) and screenshots"],
         ["HOW_TO_RUN_Step_by_Step.docx", "This guide"],
         ["Defence_Questions_and_Answers.docx", "Questions and answers for the defence"]],
        widths=[5.5, 10.5])
r.page_break()

# ---------------------------------------------------------------- Part 1
r.h1("PART 1: ONE-TIME PREPARATION")
r.h2("Step 1. Install Python")
steps(r, [
    "Open your browser and go to **https://www.python.org/downloads/**",
    "Click the big yellow button that says **Download Python 3.12** (or any number from 3.10 to 3.13).",
    "When the download finishes, open the file (it is in your Downloads folder, named like python-3.12.x-amd64.exe).",
    "On the FIRST screen of the installer, look at the bottom. Tick the box **Add python.exe to PATH**.",
    "Click **Install Now**. Wait until it says Setup was successful. Click **Close**.",
])
box(r, "WARNING: the most common mistake", [
    "If you forget to tick **Add python.exe to PATH**, nothing will work.",
    "Fix: run the installer again, choose **Modify** or **Uninstall**, then install again with the box ticked.",
], fill="FFF2CC")
r.h2("Step 2. Check that Python works")
steps(r, [
    "Press the **Windows key** on your keyboard, type **cmd**, then press **Enter**. A black window opens (Command Prompt).",
    "Type this and press Enter:",
])
typed(r, "python --version")
box(r, "You should see", ["Python 3.12.x (any version from 3.10 upwards is fine)."])
box(r, "If instead the Microsoft Store opens or you see 'python is not recognized'", [
    "Python is not on the PATH. Go back to Step 1 and reinstall with the box ticked.",
    "Then close the black window, open a new one, and try again.",
], fill="FFF2CC")
r.p("Close the black window. You will not need to type commands again, the double-click files do the work.")

r.h2("Step 3. Unzip the projects")
steps(r, [
    "Find **CMP468_Projects.zip** (usually in Downloads, or your WhatsApp or Telegram download folder).",
    "Right-click the zip file and choose **Extract All...**",
    "Click **Browse**, click **This PC**, select **Local Disk (C:)**, then click **Select Folder**.",
    "Click **Extract**. Windows creates the folder **C:\\CMP468_Projects**.",
    "Open **C:\\CMP468_Projects**. You should see the folders **project1-varsityshield**, **project2-zamanlafiya** and **reports**, plus this guide.",
])
box(r, "WARNING", [
    "Do NOT double-click files while you are still inside the zip. Windows lets you look inside a zip, "
    "but programs cannot run from there. Always extract first.",
    "If extraction made an extra folder level (C:\\CMP468_Projects\\CMP468_Projects\\project1-varsityshield), that is fine. "
    "Just open the folders until you see the 1_SETUP.bat files.",
], fill="FFF2CC")

r.h2("Step 4. Show file extensions (recommended)")
steps(r, [
    "In File Explorer, click **View** at the top (Windows 11: **View**, then **Show**).",
    "Tick **File name extensions**. Now you will see names like 1_SETUP**.bat**, which makes this guide easier to follow.",
])

r.h2("Two pop-ups you may see, and what to click")
r.table("Normal Windows warnings", ["Pop-up", "What to click"],
        [["'Windows protected your PC' (blue window) when you double-click a .bat file",
          "Click **More info**, then **Run anyway**. Windows shows this for any file downloaded from the internet."],
         ["'Windows Defender Firewall has blocked some features of Python'",
          "Tick **Private networks**, then click **Allow access**. This lets the browser talk to the program on your own laptop."]],
        widths=[7, 9])
r.page_break()

# ---------------------------------------------------------------- Part 2
r.h1("PART 2: PROJECT 1, VARSITYSHIELD")
r.p("VarsityShield watches a university server, takes encrypted backups, detects ransomware and restores "
    "the data. Open the folder **C:\\CMP468_Projects\\project1-varsityshield**.")
r.h2("Step 1. Setup (only once, internet ON)")
steps(r, [
    "Double-click **1_SETUP.bat**.",
    "A black window opens and shows steps [1/5] to [5/5]. Wait. Installing packages can take 1 to 5 minutes on slow internet.",
    "When you see **SETUP COMPLETE**, press any key to close the window.",
])
box(r, "You should see at the end", [
    "Created users 'admin' (admin) and 'auditor' (read-only) with password: ChangeMe@468",
    "Seeded university data ...",
    "Backup ...: 19 files, 19 new objects ...",
    "SETUP COMPLETE.",
])
r.h2("Step 2. Start the dashboard")
steps(r, [
    "Double-click **2_START_DASHBOARD.bat**. A black window opens. **Do not close it.** It is the program running.",
    "After a few seconds your browser opens at **http://127.0.0.1:5000**. If it shows 'This site can't be reached', wait 5 seconds and press **F5**.",
    "Log in with username **admin** and password **ChangeMe@468**.",
])
box(r, "You should see", [
    "A dark dashboard with CPU, Memory, Disk, 'Last good backup', 'Ransomware canaries: Intact'.",
    "Services table: Student Portal UP. The graph fills in over the next minute.",
])
r.figure("varsityshield_dashboard.png", "The VarsityShield dashboard", width_cm=14)
r.h2("Step 3. Run the demonstration")
steps(r, [
    "Leave the dashboard open in the browser.",
    "Go back to the project folder and double-click **3_DEMO_MENU.bat**. A menu appears.",
    "Type a number, press **Enter**, then look at the dashboard in the browser. Press any key to return to the menu.",
])
r.table("Demo order for the defence", ["Type", "What happens", "What to show and say"],
        [["1", "Student Portal is crashed", "Within 15 seconds: red alert 'Student Portal is DOWN', then it comes back UP by itself. "
               "Say: 'The system detects the outage and restarts the service automatically.'"],
         ["2", "Five result files are encrypted secretly (names unchanged)", "Nothing visible yet. Say: 'This is stealth ransomware.'"],
         ["3", "Backup is attempted", "Backup status shows **frozen** and a critical alert lists files with entropy near 8. "
               "Say: 'The encrypted files are refused, so our clean backup is protected.'"],
         ["4", "Full ransomware attack", "Canaries show **TRIPPED** in red. Type 8 to open sample_data and show the .locked files "
               "and READ_ME_TO_DECRYPT.txt."],
         ["5", "Restore from last clean backup", "Files are back. The dashboard 'Restores (measured RTO)' table shows the time. Type 8 to show the folder is clean."],
         ["6", "Verify integrity", "Prints \"ok\": true. Say: 'Every backup was decrypted and its SHA-256 checked.'"],
         ["7", "Full automatic drill", "Prints rto_seconds, rpo_seconds and byte_identical: true."]],
        widths=[1.3, 5, 9.7])
box(r, "Also show in the browser", [
    "**Backups & Restore** page: list of encrypted snapshots with Restore buttons.",
    "**Audit log** page: every action with its hash. 'Hash chain: INTACT'.",
    "Log out and log in as **auditor** (same password): the Run backup and Restore buttons are gone. That is role-based access control.",
])
r.h2("Step 4. Stop and start again later")
steps(r, [
    "To stop: close the black **2_START_DASHBOARD** window and the menu window.",
    "Next time (no internet needed): double-click **2_START_DASHBOARD.bat**, then **3_DEMO_MENU.bat**. Do NOT run 1_SETUP.bat again.",
    "To start completely fresh before the defence: close all windows, double-click **4_RESET_DEMO.bat**, type **YES**, then run **1_SETUP.bat** again.",
])
r.page_break()

# ---------------------------------------------------------------- Part 3
r.h1("PART 3: PROJECT 2, ZAMAN LAFIYA")
r.p("Zaman Lafiya is a map-based early warning system for farmer-herder conflict. Open the folder "
    "**C:\\CMP468_Projects\\project2-zamanlafiya**.")
r.h2("Step 1. Setup (only once, internet ON)")
steps(r, [
    "Double-click **1_SETUP.bat**.",
    "Wait for steps [1/6] to [6/6]. This one installs scikit-learn (a machine learning library), so it can take 3 to 8 minutes.",
    "When you see **SETUP COMPLETE**, press any key.",
])
box(r, "You should see near the end", [
    "Initialised ... 24 LGAs, zones, NDVI, contacts.",
    "Generated 1056 synthetic incidents over 104 weeks for 24 LGAs.",
    "A block of results with \"auc\" values, and \"selected\": \"random_forest\".",
    "SETUP COMPLETE.",
])
box(r, "IMPORTANT: the keys.bat file", [
    "Setup creates **keys.bat** in the folder. It holds the encryption keys. Do not delete it and do not share it.",
    "If you delete it, saved phone numbers can no longer be decrypted. Use 4_RESET_DEMO.bat to rebuild.",
], fill="FFF2CC")
r.h2("Step 2. Start the web app")
steps(r, [
    "Double-click **2_START_APP.bat**. Keep the black window open.",
    "The browser opens at **http://127.0.0.1:5050**. If not, wait 5 seconds and press **F5**.",
    "Log in with **admin** / **ChangeMe@468**.",
])
box(r, "You should see", [
    "A map with coloured circles (green, yellow, orange, red) on Nigerian LGAs, small farm squares and a dashed stock route.",
    "On the right: Highest risk LGAs, Selected LGA, Alerts dispatched, Mediation cases.",
    "With internet, a street map shows behind. Without internet the background is plain, but everything else works.",
])
r.figure("zamanlafiya_map.png", "The Zaman Lafiya live map", width_cm=15)
r.h2("Step 3. Start the live simulation")
steps(r, [
    "Keep the browser on the map.",
    "Double-click **3_START_SIMULATION.bat**. It runs for about 80 seconds.",
    "Watch the map: cow icons move along the stock route, new red dots (reports) appear, and alerts appear on the right.",
])
r.table("Demo order for the defence", ["Do this", "What to show and say"],
        [["Click the Guma circle", "The Selected LGA box shows the score, the reasons (drivers) and a trend line. "
                                   "Say: 'The system explains why an area is high risk.'"],
         ["Wait for a geofence alert", "An alert 'herd ... entered Guma yam farms' appears in English, Hausa and Pidgin, sent to "
                                       "the peace committee, the Ardo, the farmers' representative and the police."],
         ["Open Incidents (top menu)", "Click **Confirm** or **False** on a report. Point to 'protected (encrypted)' for the reporter. "
                                       "Click **Reveal** to show only an admin can see the number, and it is logged."],
         ["Open Mediation cases", "A case was opened automatically for the attack. Choose a status, type a mediator, click **Save**."],
         ["Open Audit", "Shows logins, verifications and the Reveal. 'Hash chain INTACT'."],
         ["Log out, open Public report form", "Anyone can report without logging in, with an 'anonymous' option."]],
        widths=[5, 11])
r.h2("Step 4. Stop and start again later")
steps(r, [
    "To stop: close the black windows.",
    "Next time: double-click **2_START_APP.bat**, then **3_START_SIMULATION.bat**. Do NOT run 1_SETUP.bat again.",
    "To start fresh: close all windows, double-click **4_RESET_DEMO.bat**, type **YES**.",
])
r.page_break()

# ---------------------------------------------------------------- Part 4
r.h1("PART 4: DEFENCE DAY CHECKLIST")
steps(r, [
    "The day before, run both demos fully once on the laptop you will use.",
    "Run **4_RESET_DEMO.bat** then **1_SETUP.bat** for Project 1 so the dashboard starts clean. No internet is needed because the packages are already installed.",
    "Charge the laptop fully and carry the charger.",
    "Record a short screen video of each demo (Windows key + G opens the Xbox Game Bar recorder). If anything fails on the day, play the video.",
    "Open both browser tabs and log in BEFORE you are called, so you do not type passwords in front of the panel.",
    "For the Project 2 street map, connect to a phone hotspot. The demo still works without it.",
    "Keep the reports open in Word in case you are asked to show a chapter.",
])
r.page_break()

# ---------------------------------------------------------------- Part 5
r.h1("PART 5: TROUBLESHOOTING")
r.table("Problems and fixes", ["What you see", "What to do"],
        [["'Python was not found' or 'python is not recognized'",
          "Reinstall Python with **Add python.exe to PATH** ticked (Part 1, Step 1). Close all black windows and try again."],
         ["Setup stops at 'Installing packages' with red error text about connection",
          "No internet or very weak internet. Connect to better internet and double-click 1_SETUP.bat again. It continues from where it stopped."],
         ["'Setup has not been done yet'", "Run **1_SETUP.bat** first and wait for SETUP COMPLETE."],
         ["Browser says 'This site can't be reached'",
          "Check the black 2_START window is still open. Wait 5 seconds and press F5. Type the address exactly: 127.0.0.1:5000 (Project 1) or 127.0.0.1:5050 (Project 2)."],
         ["Black window shows 'Address already in use' or closes immediately",
          "An old copy is still running. Restart the laptop, then start again."],
         ["'Invalid username or password'", "Username is **admin** (small letters), password is **ChangeMe@468** (capital C and M)."],
         ["'Too many failed attempts'", "This is the security lockout. Wait 5 minutes."],
         ["Project 1: 'Wrong backup passphrase'", "env.bat was edited after setup. Run 4_RESET_DEMO.bat, then 1_SETUP.bat."],
         ["Project 2 map has no street background", "No internet. Connect a hotspot and press F5. Everything else works without it."],
         ["Project 2 simulation shows 'HTTP Error 401'", "keys.bat changed after the app started. Close all windows and start 2_START_APP.bat again."],
         ["Project 2 phone numbers show '<undecryptable>'", "keys.bat was deleted or replaced. Run 4_RESET_DEMO.bat."],
         ["Antivirus deletes or blocks a file", "The demo 'ransomware' is harmless and only touches the sample_data folder. Allow the folder C:\\CMP468_Projects in your antivirus, or pause it during the demo."],
         ["Blue 'Windows protected your PC' box", "Click **More info**, then **Run anyway**."]],
        widths=[6, 10])
r.page_break()

# ---------------------------------------------------------------- Appendices
r.h1("APPENDIX A: WHAT THE DOUBLE-CLICK FILES DO")
r.p("If a lecturer asks you to run the system from the command line, these are the same steps by hand. "
    "Open the project folder, click the address bar at the top of File Explorer, type **cmd** and press Enter. "
    "A black window opens already inside the folder.")
r.h3("Project 1")
typed(r, "python -m venv venv\nvenv\\Scripts\\activate\npip install -r requirements.txt\nenv.bat\npython cli.py init\n"
         "python simulate.py seed\npython cli.py backup\npython cli.py verify\npython app.py")
r.h3("Project 2")
typed(r, "python -m venv venv\nvenv\\Scripts\\activate\npip install -r requirements.txt\npython manage.py genkey --save\n"
         "keys.bat\npython manage.py init\npython manage.py history\npython manage.py train\npython manage.py risk\npython app.py")
r.p("To send a test SMS by hand while Project 2 runs (open a second black window in the same folder):")
typed(r, 'curl.exe -X POST -d "from=08031234567&text=ZL Guma CROP cows in yam farm" http://127.0.0.1:5050/api/sms')

r.h1("APPENDIX B: MAC OR LINUX")
r.p("Install Python 3.10+ (Mac: from python.org. Ubuntu: sudo apt install python3 python3-venv). Open Terminal in "
    "the project folder and type:")
r.h3("Project 1")
typed(r, "python3 -m venv venv\nsource venv/bin/activate\npip install -r requirements.txt\n"
         "export VSHIELD_PASSPHRASE='CMP468-VarsityShield-Demo-Passphrase'\npython cli.py init\n"
         "python simulate.py seed\npython cli.py backup\npython simulate.py portal --background\npython app.py")
r.h3("Project 2")
typed(r, "python3 -m venv venv\nsource venv/bin/activate\npip install -r requirements.txt\npython manage.py genkey --save\n"
         "source keys.sh\npython manage.py init\npython manage.py history\npython manage.py train\npython app.py")
r.p("In a second Terminal, activate the venv again (and run source keys.sh for Project 2) before running the "
    "demo commands from the menus: python simulate.py outage, tamper, attack, drill, or python simulate.py --steps 40 "
    "--interval 2 for Project 2.")
r.save(os.path.join(HERE, "HOW_TO_RUN_Step_by_Step.docx"))
print("wrote run guide")

# =================================================================== DEFENCE Q&A
d = Report()
cover(d, "DEFENCE QUESTIONS AND ANSWERS", "Likely panel questions with short, clear answers")

d.h1("HOW TO USE THIS DOCUMENT")
d.p("Do not memorise the answers word for word. Panels can tell. Read each answer, understand the idea, then say it "
    "in your own words. If you do not know an answer, say what you do know and how you would find out. That is "
    "better than guessing.")
d.p("Before the defence, open each project's code and find the file listed in Section G for every feature. Be ready "
    "to scroll to it and explain it in two sentences.")

d.h1("SECTION A: OPENING STATEMENTS (ABOUT 1 MINUTE EACH)")
d.h3("Project 1: VarsityShield")
d.p("\"Nigerian universities now keep results, fees and payroll on servers, but most protect them only with "
    "antivirus and firewalls. Studies show weak incident response and recovery in our universities. If ransomware "
    "strikes, a whole session's results can be lost. VarsityShield monitors servers and services, takes encrypted "
    "backups in three places, stops ransomware-encrypted files from entering the backups, and restores data with "
    "measured recovery time. In my tests it detected a simulated attack and restored all files exactly, in under "
    "a second for the demo data. I will now demonstrate it.\"")
d.h3("Project 2: Zaman Lafiya")
d.p("\"Farmer-herder conflict kills people and increases food insecurity in the Middle Belt. Warnings come late, "
    "and people who report incidents can become targets. Zaman Lafiya is a GIS early warning system. It maps farms, "
    "grazing reserves and stock routes, accepts reports by web, SMS and USSD, tracks herds by GPS, gives every LGA "
    "an explained risk score, and sends alerts in English, Hausa and Pidgin to local leaders. Reporters' phone "
    "numbers are encrypted. I will now show it working.\"")

d.h1("SECTION B: GENERAL COMPUTER SECURITY QUESTIONS (CMP 468)")
qa(d, [
    ("What is the CIA triad and how do your projects use it?",
     "Confidentiality, Integrity, Availability. In VarsityShield, AES-256-GCM encryption gives confidentiality, "
     "SHA-256 hashes and the hash chain give integrity, and three backup copies with tested restores give "
     "availability. In Zaman Lafiya, encrypted phone numbers give confidentiality, signed GPS messages and the audit "
     "chain give integrity, and the offline map view and SMS/USSD channels support availability."),
    ("What is the difference between a threat, a vulnerability and a risk?",
     "A **threat** is something that can cause harm, for example ransomware. A **vulnerability** is a weakness it can "
     "use, for example a backup stored on a network share with no encryption. **Risk** is the chance of harm times "
     "its impact. Controls reduce risk by removing vulnerabilities or reducing impact."),
    ("What is the difference between encryption and hashing?",
     "Encryption is two-way: with the key you get the original data back. Hashing is one-way: you cannot get the data "
     "back, you can only compare. I encrypt backups because I must restore them, and I hash them (SHA-256) to check "
     "they were not changed."),
    ("Symmetric or asymmetric encryption? Which did you use and why?",
     "Symmetric: the same key encrypts and decrypts (AES). Asymmetric uses a public and a private key (RSA). I used "
     "symmetric AES because it is fast for large data and only the university itself needs to decrypt. Asymmetric "
     "would help if many parties had to send encrypted data to one receiver."),
    ("What is AES-256-GCM?",
     "AES is the Advanced Encryption Standard. 256 is the key length in bits. GCM (Galois/Counter Mode) encrypts and "
     "also produces an authentication tag. If one byte of the backup is changed, decryption fails, so tampering is "
     "detected automatically."),
    ("What is a salt and why do you use it?",
     "A random value added before hashing a password or deriving a key. Two users with the same password get "
     "different hashes, and attackers cannot use precomputed tables. Both projects store a random salt with each "
     "password hash."),
    ("Why not store passwords with plain SHA-256 or MD5?",
     "They are too fast, so an attacker can try billions of guesses per second. MD5 is also broken. I used PBKDF2 "
     "with 310,000 iterations, which slows each guess down on purpose."),
    ("What is an HMAC?",
     "A hash computed with a secret key. Only someone with the key can produce the right value. Zaman Lafiya uses it "
     "to sign GPS messages and to create pseudonyms for reporters. VarsityShield uses it to name backup objects."),
    ("What is ransomware and how does it usually get in?",
     "Malware that encrypts files and demands payment for the key. It usually enters through phishing emails, weak "
     "remote desktop passwords, or unpatched software. Modern ransomware also looks for backups to destroy them, "
     "which is why my backups are encrypted, read-only, replicated and checked."),
    ("What is defence in depth?",
     "Using several layers of control so one failure does not lead to disaster. In VarsityShield: monitoring, "
     "canaries, entropy checks, encryption, replicas and audit logs each cover a different failure."),
    ("What are preventive, detective and corrective controls? Give examples from your work.",
     "Preventive stops an attack: login lockout, encryption, role-based access. Detective finds it: monitoring, canary "
     "files, integrity verification. Corrective fixes the damage: restore, repair from replica, automatic service "
     "restart."),
    ("What is role-based access control (RBAC)?",
     "Users get permissions through roles, not one by one. VarsityShield has admin and viewer. Zaman Lafiya has admin, "
     "analyst, responder and viewer. A responder can update cases but cannot see informant numbers."),
    ("What is CSRF and how did you prevent it?",
     "Cross-Site Request Forgery tricks a logged-in user's browser into sending a request they did not intend, for "
     "example 'restore backup'. Every form carries a random secret token tied to the session. A request without the "
     "right token is rejected with error 400."),
    ("How did you prevent SQL injection?",
     "All database queries use parameters (question marks) instead of joining user text into SQL. The database "
     "treats user input as data, never as commands."),
    ("What is the Nigeria Data Protection Act 2023 and how does it affect your projects?",
     "It is the law on personal data in Nigeria, enforced by the Nigeria Data Protection Commission. It requires "
     "data to be kept secure, accurate and only as long as needed. My projects encrypt personal data, restrict who "
     "can see it, log access, and recommend retention limits."),
    ("Which security standards did you follow?",
     "NIST Cybersecurity Framework 2.0 (especially the Recover function), ISO/IEC 27001:2022 (for example control "
     "8.13, information backup), and the NDPA 2023."),
])

d.h1("SECTION C: VARSITYSHIELD QUESTIONS")
qa(d, [
    ("What problem does VarsityShield solve?",
     "Nigerian universities often have no monitoring, untested backups, and backups that ransomware can reach. "
     "VarsityShield gives continuous monitoring, encrypted 3-2-1 backups, ransomware detection and tested recovery."),
    ("What are RPO and RTO? What values did you use?",
     "RPO (Recovery Point Objective) is how much data, in time, you can afford to lose. RTO (Recovery Time Objective) "
     "is how long recovery may take. The policy uses RPO of 60 minutes with backups every 30 minutes. In the drill, "
     "RTO was about 0.15 seconds for 9 MB of demo data."),
    ("What is the 3-2-1 backup rule?",
     "Three copies of data, on two different media, with one copy offsite. VarsityShield writes the main repository, "
     "a second-disk (NAS) replica and an offsite replica. The extended 3-2-1-1-0 rule adds one offline or immutable "
     "copy and zero errors after verification."),
    ("How does VarsityShield detect ransomware?",
     "Two ways. First, canary files: decoy files nobody should touch. If their hash changes, something is mass-editing "
     "files. Second, every changed file is checked before backup: text files whose entropy jumps near 8 bits per byte, "
     "or Word/PDF files whose signature bytes disappear, are treated as encrypted."),
    ("What is entropy?",
     "A measure of randomness from 0 to 8 bits per byte. Normal text is around 4 to 5. Encrypted data is close to 8 "
     "because it looks random. In the demo the encrypted CSV files measured 7.95 to 7.99."),
    ("What does 'backup frozen' mean?",
     "When the system suspects ransomware it refuses to create a new snapshot, so encrypted files never replace the last "
     "clean version. The clean history stays available for restore."),
    ("Why do you not name backup objects with a normal SHA-256 hash?",
     "Anyone who saw the repository could hash a known file, for example a leaked payroll, and check whether it is "
     "stored. Using HMAC with a secret key gives the same de-duplication without that leak."),
    ("What is de-duplication and incremental backup?",
     "Incremental means only changed files are read and stored. De-duplication means identical content is stored "
     "once. In the test, a backup with no changes took 0.09 seconds and stored nothing new."),
    ("How do you know a backup has not been secretly changed?",
     "Three checks. The AES-GCM tag fails if any byte changes. The SHA-256 of every file is compared after "
     "decryption. Each snapshot stores the hash of the previous one, so deleting or editing an old snapshot breaks "
     "the chain. In my test, changing one byte was detected and the repair command fixed it from a replica."),
    ("What happens if the administrator forgets the passphrase?",
     "The backups cannot be decrypted. That is the price of strong encryption. The policy recommends sealing a copy "
     "in two physical locations. Future work is splitting it among three officers with secret sharing."),
    ("Can ransomware on the server delete the backups too?",
     "It can try to delete the local repository if it gets administrator rights. That is why one replica must be "
     "offsite or offline, such as a rotated USB disk or an object-locked cloud bucket. Objects are also read-only."),
    ("How does monitoring work and what is the anomaly detection?",
     "Every 15 seconds it records CPU, memory, disk, network and battery, and checks services. Besides fixed limits "
     "(e.g. CPU above 85%), it compares each reading to the last 30: if it is more than 3 standard deviations away "
     "(a z-score above 3), it raises an anomaly alert."),
    ("Why monitor battery? That is unusual.",
     "Power is unstable in Nigeria, so servers run on inverters and UPS. A server on low battery is about to go down. "
     "The alert gives the ICT team time to run an emergency backup and shut down safely."),
    ("What is the audit log and why is it tamper-evident?",
     "It records every login, backup, restore and verification. Each entry contains the hash of the previous entry. "
     "If someone deletes or changes a row, the chain no longer matches and the audit page shows BROKEN."),
    ("Your test data is only 9 MB. Is the result realistic?",
     "The timing shows that the method works, not production speed. For a real 200 GB server, restore time depends "
     "on disk and network speed, so the report recommends measuring RTO in quarterly drills with real data."),
    ("How is this different from just copying files to an external drive?",
     "A manual copy is unencrypted, untested, often forgotten, sits in the same room, and ransomware can encrypt it. "
     "VarsityShield is automatic, encrypted, verified, replicated, ransomware-aware, and keeps evidence of every action."),
    ("How would a university deploy it?",
     "Install it on a backup server, point the sources at database dumps and file folders, set replicas to a NAS "
     "and an offsite location, enable email or SMS alerts, and run it as a Windows service or Linux systemd service. "
     "The README has the steps."),
])

d.h1("SECTION D: ZAMAN LAFIYA QUESTIONS")
qa(d, [
    ("What does Zaman Lafiya mean and what does the system do?",
     "It is Hausa for 'living in peace'. The system maps farms, grazing reserves and stock routes, collects incident "
     "reports, tracks herds, scores risk for each LGA, sends local-language alerts and tracks mediation cases."),
    ("What is GIS and why use it here?",
     "A Geographic Information System stores and analyses data by location. Farmer-herder conflict happens at "
     "specific places and seasons: farms near stock routes, dry areas, river crossings. GIS lets us see and calculate "
     "those overlaps."),
    ("How is the risk score calculated?",
     "Two parts. A rule-based index adds weighted factors: recent incidents, fatalities, vegetation below normal, crop "
     "season, herd movement season, rumours, history, and herds near or inside farms. A random forest model also "
     "predicts the chance of an incident in the next 7 days. The final score is the average of both, from 0 to 100: "
     "Low, Moderate, High or Severe."),
    ("What is NDVI and why is it important?",
     "Normalized Difference Vegetation Index, measured by satellite, shows how green vegetation is. When pasture is "
     "drier than normal, herds move further into farm areas, which raises conflict risk. In my model, the NDVI "
     "anomaly was the most important feature (0.317)."),
    ("What is a geofence?",
     "A virtual boundary on a map. When a herd's GPS position falls inside a farm polygon, or more than 15 km away "
     "from the stock route, the system sends an alert immediately, before damage turns into violence."),
    ("How do you know a point is inside a farm?",
     "The ray-casting algorithm: draw an imaginary line from the point and count how many polygon edges it crosses. "
     "Odd means inside, even means outside."),
    ("Your model AUC is 0.659. Isn't that low?",
     "Predicting an incident in a specific LGA within 7 days is hard. AUC of 0.5 is guessing, and the simple "
     "'last month's count' method scored 0.611, so the model adds value. The rule layer adds live signals such as herds "
     "inside farms. With real ACLED data and real NDVI the model should be retrained and re-evaluated."),
    ("What is AUC?",
     "Area Under the ROC Curve. It measures how well the model ranks risky weeks above safe weeks. 1.0 is perfect, 0.5 "
     "is random. I used it because accuracy is misleading when most weeks have no incident."),
    ("Your data is synthetic. Why?",
     "Real incident data with exact locations is sensitive and ACLED requires registration. Synthetic data with "
     "realistic seasons and drought effects lets me test the full pipeline. The system has an import command for "
     "real ACLED data, and the report states this limitation clearly."),
    ("Why did you split training and testing by time?",
     "So the model is tested only on the future, like real use. A random split would let it 'see' future weeks "
     "during training and give an unfairly high score."),
    ("How can people without smartphones report?",
     "By SMS, for example 'ZL Guma CROP cows in yam farm', or by USSD, a dial menu that works on any phone. The web "
     "form is for smartphones."),
    ("How do you protect people who report incidents?",
     "Phone numbers are encrypted with Fernet (AES with HMAC). For counting and credibility, the system uses a keyed "
     "hash (pseudonym) instead of the number. Only an admin can reveal a number, must give a reason, and the reveal is "
     "written to the audit log. People can also report anonymously."),
    ("How do you stop fake reports or rumours from causing violence?",
     "Reports are verified by an analyst before escalation. Unverified reports count less in the score. Each sender "
     "gets a credibility score from past confirmed and false reports. There is also a limit of 5 reports per number "
     "per hour, and a hidden honeypot field blocks bots on the web form."),
    ("Can someone send fake herd locations?",
     "Each GPS message must carry a signature: HMAC-SHA256 of the timestamp and the message, made with a secret key. "
     "Messages without the right signature, or older than 5 minutes (replay), are rejected with error 401."),
    ("Why send alerts in Hausa and Pidgin?",
     "Many herders and rural farmers do not read English well. An alert they cannot read is useless. Local language "
     "also shows respect and builds trust in the system."),
    ("Will herders accept being tracked?",
     "It must be voluntary and agreed with herders' associations. Herders also benefit: route guidance, water point "
     "status, and evidence that protects them from false accusations."),
    ("Does the system target or stereotype an ethnic group?",
     "No. Reports describe events, such as crop destruction or cattle theft, not ethnicity. Alerts go to both farmer "
     "and herder leaders at the same time, and verification is required before escalation."),
    ("How does a warning lead to action?",
     "Every serious report automatically opens a mediation case. Responders record the mediator, status and any "
     "compensation until the case is resolved or escalated. This closes the gap between warning and response that "
     "studies found in existing systems."),
    ("Why SQLite instead of PostGIS?",
     "SQLite needs no installation, which suits a prototype on a student laptop. For statewide use, PostGIS is the "
     "recommended upgrade."),
    ("Why is loading a saved model file a security issue?",
     "Python pickle files can run code when loaded. A tampered model file could take over the server. Zaman Lafiya "
     "saves the SHA-256 hash at training time and refuses to load a file whose hash does not match."),
])

d.h1("SECTION E: METHODOLOGY AND REPORT QUESTIONS")
qa(d, [
    ("What research methodology did you use?",
     "Design Science Research: identify the problem from literature, set objectives, design and build the system, "
     "demonstrate it, evaluate it with measurable tests, and report the results."),
    ("Why only literature from 2021 to 2026?",
     "To reflect current threats and technologies. Ransomware tactics, monitoring tools and climate-conflict "
     "research have changed a lot in the last five years."),
    ("What is the research gap?",
     "Project 1: no reviewed work combines monitoring, encrypted 3-2-1 backup, ransomware gating, verification and "
     "audited recovery in a low-cost tool for Nigerian universities. Project 2: no reviewed system for Nigeria joins "
     "GIS layers, basic-phone reporting, herd geofencing, explainable risk scoring, local-language alerts, mediation "
     "tracking and informant protection."),
    ("What are the limitations of your work?",
     "Project 1: small test data, entropy cannot judge already-compressed files, and an offline copy is needed for full "
     "protection. Project 2: synthetic training data, approximate coordinates, and tracking depends on herders' consent."),
    ("What is your contribution to knowledge?",
     "Project 1: an open design joining monitoring, encrypted 3-2-1 backup, ransomware gating and audited recovery, "
     "plus HMAC object names that de-duplicate without leaking information. Project 2: an integrated, secure early "
     "warning design for Nigeria that protects informants by design."),
    ("What would you do with more time?",
     "Project 1: direct database backup with point-in-time recovery, Prometheus/Grafana export, secret sharing of the "
     "passphrase. Project 2: train on real ACLED and MODIS data, Hausa and Pidgin text classification, an offline "
     "Android app, and satellite crop maps."),
    ("Which tools and languages did you use?",
     "Python, Flask for the web interface, SQLite for storage, the cryptography library for AES and Fernet, psutil "
     "for system metrics, scikit-learn for machine learning, and Leaflet with OpenStreetMap for the map."),
])

d.h1("SECTION F: DIFFICULT QUESTIONS")
qa(d, [
    ("Did you write all this code yourself?",
     "Answer truthfully, and follow your department's rules on tools and help. If you used AI assistance or "
     "tutorials, say so plainly and show that you understand every part: open any file the panel names and explain "
     "what it does. Panels usually judge understanding more than typing."),
    ("Show me where the encryption happens in the code.",
     "Project 1: open backup_engine.py and find the functions _seal (encrypts) and _open (decrypts). Project 2: open "
     "core.py and find encrypt_pii and decrypt_pii."),
    ("What happens if I unplug the internet now?",
     "Both systems keep working because they run on the laptop. Only the street map background in Project 2 "
     "disappears. Alerts are stored and shown, and in a real deployment the SMS gateway would need a connection."),
    ("What if your system itself is hacked?",
     "Attackers would face login lockout, hashed passwords, CSRF tokens and role limits. Personal data and backups "
     "are encrypted, so a stolen database or backup is unreadable without the keys. The audit chain shows tampering. "
     "Keys are kept outside the database, in environment variables."),
    ("Why should a university pay for this?",
     "It costs nothing in licences. It runs on existing hardware with free software. The cost is staff time for "
     "setup and quarterly drills, which is far less than losing a session's results or paying a ransom."),
])

d.h1("SECTION G: WHERE THINGS ARE IN THE CODE")
d.table("Project 1: VarsityShield", ["Feature", "File", "Function or place"],
        [["Encryption / decryption", "backup_engine.py", "_seal, _open"],
         ["Key from passphrase (scrypt)", "backup_engine.py", "_derive, load_keys"],
         ["Incremental backup", "backup_engine.py", "run_backup"],
         ["Verify and repair", "backup_engine.py", "verify_repository, repair_from_replicas"],
         ["Restore", "backup_engine.py", "restore"],
         ["Retention policy", "backup_engine.py", "prune"],
         ["Canary files and entropy", "ransomware_guard.py", "check_canaries, inspect_file, shannon_entropy"],
         ["Monitoring and anomaly (z-score)", "monitor.py", "run_once, _zscore_anomaly"],
         ["Automatic service restart", "monitor.py", "_run_remediation"],
         ["Alerts (email, Telegram, SMS)", "common.py", "raise_alert"],
         ["Password hashing", "common.py", "hash_password, check_password"],
         ["Audit hash chain", "common.py", "audit, verify_audit_chain"],
         ["Login, roles, CSRF", "app.py", "login, login_required"],
         ["Demo attack and drill", "simulate.py", "attack, tamper_in_place, drill"]],
        widths=[5.5, 4, 6.5])
d.table("Project 2: Zaman Lafiya", ["Feature", "File", "Function or place"],
        [["LGAs, farms, routes", "geodata.py", "LGAS, STOCK_ROUTE, build_zones"],
         ["Phone encryption", "core.py", "encrypt_pii, decrypt_pii"],
         ["Reporter pseudonym", "core.py", "pseudonym"],
         ["Point in polygon / distance", "core.py", "point_in_polygon, distance_to_line_km"],
         ["Risk features", "core.py", "lga_features, herd_pressure"],
         ["Rule-based score", "core.py", "rule_score"],
         ["Final score with ML", "core.py", "compute_risk, load_model"],
         ["Geofence alerts", "core.py", "check_geofence"],
         ["Alert messages (3 languages)", "core.py", "TEMPLATES, ENCROACH, dispatch"],
         ["SMS parsing", "core.py", "parse_sms"],
         ["Credibility and rate limit", "core.py", "add_incident, reporter_credibility"],
         ["Model training and evaluation", "manage.py", "build_dataset, cmd_train"],
         ["ACLED import", "manage.py", "cmd_import_acled"],
         ["SMS, USSD, GPS endpoints", "app.py", "api_sms, api_ussd, herd_ping"],
         ["Roles and CSRF", "app.py", "require"],
         ["Reveal informant (logged)", "app.py", "reveal"]],
        widths=[5.5, 3.5, 7])
d.save(os.path.join(HERE, "Defence_Questions_and_Answers.docx"))
print("wrote defence Q&A")
